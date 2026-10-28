#!/usr/bin/env python3
"""Generate browser reel/stop data from the approved NOVA Google Sheet export."""

from __future__ import annotations

import hashlib
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


SOURCE_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "16wF24WmveUjq2mBubo593AJC40Kx7BODwha8S2TYeNc/edit?gid=310238419"
)
SYMBOLS = {
    "リプ": "REPLAY",
    "ベル": "🔔",
    "赤７": "7",
    "チェリー": "🍒",
    "スイカ": "🍋",
    "ネビュラ": "NEBULA",
    "ブランク": "__BLANK__",
    "BAR": "BAR",
}
ROLE_PREFIXES = (
    ("BAR菱形揃いの場合の停止形", "SUPER_NOVA"),
    ("ネビュラ揃いの場合の停止形", "NEBULA"),
    ("BAR揃いの場合の停止形", "BAR3"),
    ("赤７揃いの場合の停止形", "BIG"),
    ("外れ目の停止形", "MISS"),
    ("リプレイの停止形", "REPLAY"),
    ("８枚ベルの停止形", "BELL"),
    ("１５枚ベルの停止形", "BELL15"),
    ("弱チェリーの停止形", "WEAK_NOVA"),
    ("強チェリーの停止形", "STRONG_NOVA"),
    ("スイカの停止形", "WEAK_SUICA"),
)
EXPECTED_COUNTS = {
    "MISS": 175,
    "REPLAY": 64,
    "BELL": 64,
    "BELL15": 64,
    "WEAK_NOVA": 64,
    "STRONG_NOVA": 48,
    "WEAK_SUICA": 36,
    "NEBULA": 1,
    "BAR3": 1,
    "BIG": 4,
    "SUPER_NOVA": 1,
}


def role_for_title(title: str) -> str | None:
    for prefix, role in ROLE_PREFIXES:
        if title.startswith(prefix):
            return role
    return None


def convert(value: object, *, where: str) -> str:
    try:
        return SYMBOLS[str(value)]
    except KeyError as exc:
        raise ValueError(f"Unknown symbol at {where}: {value!r}") from exc


def extract_reels(sheet: openpyxl.worksheet.worksheet.Worksheet) -> list[list[str]]:
    reels = [[], [], []]
    for row in range(4, 24):
        expected = row - 3
        actual = sheet.cell(row, 2).value
        if int(actual) != expected:
            raise ValueError(f"Unexpected reel index at B{row}: {actual!r}")
        for column in range(3, 6):
            reels[column - 3].append(
                convert(sheet.cell(row, column).value, where=sheet.cell(row, column).coordinate)
            )
    if any(len(reel) != 20 for reel in reels):
        raise ValueError("Each reel must contain exactly 20 symbols")
    return reels


def extract_stops(sheet: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, list[list[list[str]]]]:
    stops: dict[str, list[list[list[str]]]] = defaultdict(list)
    for row in sheet.iter_rows():
        for cell in row:
            title = cell.value
            if not isinstance(title, str) or "停止形" not in title or title.startswith("※"):
                continue
            role = role_for_title(title)
            if not role:
                continue
            raw_matrix = [
                [
                    sheet.cell(cell.row + row_offset, cell.column + column_offset).value
                    for column_offset in range(3)
                ]
                for row_offset in range(1, 4)
            ]
            # Category headings also contain the word 「停止形」. Only the named
            # pattern headings have a complete 3x3 matrix immediately below them.
            if any(value is None for row_values in raw_matrix for value in row_values):
                continue
            matrix = []
            for row_offset in range(1, 4):
                matrix.append(
                    [
                        convert(
                            raw_matrix[row_offset - 1][column_offset],
                            where=sheet.cell(cell.row + row_offset, cell.column + column_offset).coordinate,
                        )
                        for column_offset in range(3)
                    ]
                )
            stops[role].append(matrix)
    counts = {role: len(values) for role, values in stops.items()}
    if counts != EXPECTED_COUNTS:
        raise ValueError(f"Unexpected stop counts: {counts!r}")
    return dict(stops)


def verify_contiguous(reels: list[list[str]], stops: dict[str, list[list[list[str]]]]) -> None:
    for role, matrices in stops.items():
        for pattern_index, matrix in enumerate(matrices):
            for reel_index, reel in enumerate(reels):
                column = [matrix[row][reel_index] for row in range(3)]
                if not any(
                    all(column[offset] == reel[(top + offset) % len(reel)] for offset in range(3))
                    for top in range(len(reel))
                ):
                    raise ValueError(
                        f"{role}[{pattern_index}] reel {reel_index} is not a contiguous three-symbol window"
                    )


def main() -> None:
    if len(sys.argv) not in (2, 3):
        raise SystemExit("usage: generate-nova-reel-sheet-data.py SOURCE.xlsx [OUTPUT.js]")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve() if len(sys.argv) == 3 else Path("nova-reel-sheet-data.js").resolve()
    # The exported workbook has a large declared sheet range; regular mode is
    # substantially faster here than streaming every empty cell in that range.
    book = openpyxl.load_workbook(source, read_only=False, data_only=True)
    reels = extract_reels(book["リール配列"])
    stops = extract_stops(book["停止形一覧"])
    verify_contiguous(reels, stops)
    sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    data = {
        "version": "google-sheet-20261027",
        "sourceUrl": SOURCE_URL,
        "sourceSha256": sha256,
        "reelsRange": "リール配列!C4:E23",
        "stopsRange": "停止形一覧!A1:FP204",
        "patternCount": sum(map(len, stops.values())),
        "counts": dict(Counter({role: len(values) for role, values in stops.items()})),
        "strips": reels,
        "stops": stops,
    }
    encoded = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    output.write_text(
        "// Generated by scripts/generate-nova-reel-sheet-data.py. Do not edit by hand.\n"
        "// Source: " + SOURCE_URL + "\n"
        "globalThis.NovaReelSheetData=(()=>{\n"
        " const deepFreeze=value=>{if(value&&typeof value==='object'&&!Object.isFrozen(value)){Object.freeze(value);for(const item of Object.values(value))deepFreeze(item);}return value;};\n"
        " return deepFreeze(" + encoded + ");\n"
        "})();\n",
        encoding="utf-8",
        newline="\n",
    )
    print(f"Generated {output} with {data['patternCount']} stop patterns ({sha256})")


if __name__ == "__main__":
    main()
